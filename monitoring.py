#!/usr/bin/env python3

import configparser
import os
import re
from datetime import datetime, timedelta
from pathlib import Path
from urllib.parse import urlparse

import paramiko
import requests
import urllib3

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ============================================================
# PATHS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = BASE_DIR / "monitoring_output"
REPORT_FILE = OUTPUT_DIR / "monitoring_report.xlsx"

URL_FILE = BASE_DIR / "urls.txt"
MONGO_FILE = BASE_DIR / "mongodb.txt"
RABBITMQ_FILE = BASE_DIR / "rabbitmq.txt"
CASSANDRA_FILE = BASE_DIR / "cassandra.txt"
CREDENTIALS_FILE = BASE_DIR / "credentials.txt"


# ============================================================
# GENERAL SETTINGS
# ============================================================

URL_TIMEOUT = 10
SSH_TIMEOUT = 15
COMMAND_TIMEOUT = 120

SUCCESS_HTTP_CODES = {
    200, 201, 202, 204,
    301, 302, 303, 307, 308
}

# Internal/self-signed HTTPS certificates are intentionally allowed.
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)


# ============================================================
# CONFIG FILE HELPERS
# ============================================================

def load_key_value_file(path):
    """
    Reads files like:

    [SECTION]
    key=value
    key=value

    Returns:
        {
            "SECTION": {
                "key": "value"
            }
        }
    """
    if not path.exists():
        raise FileNotFoundError(f"Configuration file not found: {path}")

    parser = configparser.ConfigParser(
        interpolation=None,
        delimiters=("="),
        inline_comment_prefixes=None,
        strict=False
    )

    parser.optionxform = str
    parser.read(path, encoding="utf-8")

    result = {}

    for section in parser.sections():
        result[section] = dict(parser.items(section))

    return result


def load_urls():
    """
    URL file format:

    [NAVFAS]
    Service Name=http://example/

    [NAVRPS]
    Service Name=http://example/
    """

    if not URL_FILE.exists():
        raise FileNotFoundError(f"URL file not found: {URL_FILE}")

    urls = []
    current_project = None

    with open(URL_FILE, "r", encoding="utf-8") as file:
        for line_number, raw_line in enumerate(file, start=1):

            line = raw_line.strip()

            if not line:
                continue

            if line.startswith("#"):
                continue

            if line.startswith("[") and line.endswith("]"):
                current_project = line[1:-1].strip()

                if not current_project:
                    raise ValueError(
                        f"Empty project name at {URL_FILE}:{line_number}"
                    )

                continue

            if "=" not in line:
                print(
                    f"WARNING: Ignoring invalid URL line "
                    f"{URL_FILE}:{line_number}: {line}"
                )
                continue

            if not current_project:
                print(
                    f"WARNING: URL found before a project section "
                    f"at {URL_FILE}:{line_number}"
                )
                continue

            service, url = line.split("=", 1)
            service = service.strip()
            url = url.strip()

            if not service:
                print(
                    f"WARNING: Empty service name at "
                    f"{URL_FILE}:{line_number}"
                )
                continue

            if not url:
                print(
                    f"WARNING: Empty URL for {service} at "
                    f"{URL_FILE}:{line_number}"
                )
                continue

            if not url.lower().startswith(("http://", "https://")):
                print(
                    f"WARNING: Ignoring non-HTTP URL at "
                    f"{URL_FILE}:{line_number}: {url}"
                )
                continue

            urls.append({
                "project": current_project,
                "service": service,
                "url": url
            })

    return urls


def load_credentials():
    config = load_key_value_file(CREDENTIALS_FILE)

    if "SSH" not in config:
        raise ValueError(
            f"[SSH] section missing in {CREDENTIALS_FILE}"
        )

    username = config["SSH"].get("username", "").strip()
    password = config["SSH"].get("password", "")

    if not username:
        raise ValueError("SSH username is empty in credentials.txt")

    if not password:
        raise ValueError("SSH password is empty in credentials.txt")

    return username, password


# ============================================================
# DATE
# ============================================================

def get_backup_date():
    """
    Always returns yesterday in d%bY format.

    Example:
        Today     = 22Sep2026
        Backup    = 21Sep2026
    """
    yesterday = datetime.now() - timedelta(days=1)
    return yesterday.strftime("%d%b%Y")


def was_yesterday_sunday():
    yesterday = datetime.now() - timedelta(days=1)
    return yesterday.weekday() == 6


# ============================================================
# SSH
# ============================================================

SSH_USER, SSH_PASSWORD = load_credentials()


def ssh_connect(server):
    client = paramiko.SSHClient()
    client.set_missing_host_key_policy(paramiko.AutoAddPolicy())

    try:
        client.connect(
            hostname=server,
            username=SSH_USER,
            password=SSH_PASSWORD,
            timeout=SSH_TIMEOUT,
            auth_timeout=SSH_TIMEOUT,
            banner_timeout=SSH_TIMEOUT,
            look_for_keys=False,
            allow_agent=False
        )

        return client

    except Exception as exc:
        print(f"SSH connection failed: {server}")
        print(f"Reason: {exc}")
        return None


def execute_ssh_command(server, command, use_sudo=False):
    """
    Executes a command over SSH.

    If use_sudo=True, sudo receives the password through stdin.
    The password is NOT placed in the shell command itself.
    """

    client = None

    try:
        print(f"SSH: {server}")
        print(f"CMD: {command}")

        client = ssh_connect(server)

        if client is None:
            return {
                "success": False,
                "output": "",
                "error": "SSH connection failed",
                "exit_code": -1
            }

        if use_sudo:
            remote_command = f"sudo -S -p '' {command}"
        else:
            remote_command = command

        stdin, stdout, stderr = client.exec_command(
            remote_command,
            timeout=COMMAND_TIMEOUT
        )

        if use_sudo:
            stdin.write(SSH_PASSWORD + "\n")
            stdin.flush()

        output = stdout.read().decode(
            "utf-8",
            errors="replace"
        ).strip()

        error = stderr.read().decode(
            "utf-8",
            errors="replace"
        ).strip()

        exit_code = stdout.channel.recv_exit_status()

        # sudo may print the password prompt or warnings to stderr.
        # We only treat non-zero exit code as command failure.
        if exit_code != 0:
            print(f"SSH COMMAND FAILED: {server}")
            if error:
                print(f"Reason: {error}")

        return {
            "success": exit_code == 0,
            "output": output,
            "error": error,
            "exit_code": exit_code
        }

    except Exception as exc:
        print(f"SSH command execution failed: {server}")
        print(f"Reason: {exc}")

        return {
            "success": False,
            "output": "",
            "error": str(exc),
            "exit_code": -1
        }

    finally:
        if client:
            client.close()


# ============================================================
# URL CHECK
# ============================================================

def alternate_url(url):
    parsed = urlparse(url)

    if parsed.scheme.lower() == "https":
        return url.replace("https://", "http://", 1)

    if parsed.scheme.lower() == "http":
        return url.replace("http://", "https://", 1)

    return None


def try_url(url):
    try:
        response = requests.get(
            url,
            timeout=URL_TIMEOUT,
            verify=False,
            allow_redirects=True
        )

        if response.status_code in SUCCESS_HTTP_CODES:
            return {
                "success": True,
                "working_url": response.url,
                "http_code": response.status_code,
                "error": ""
            }

        return {
            "success": False,
            "working_url": "",
            "http_code": response.status_code,
            "error": f"HTTP {response.status_code}"
        }

    except requests.RequestException as exc:
        return {
            "success": False,
            "working_url": "",
            "http_code": "",
            "error": str(exc)
        }


def check_url(url):
    # First attempt: configured URL
    first = try_url(url)

    if first["success"]:
        return {
            "status": "Yes",
            "working_url": first["working_url"],
            "http_code": first["http_code"],
            "details": ""
        }

    # Second attempt: HTTP <-> HTTPS
    alt = alternate_url(url)

    if alt:
        second = try_url(alt)

        if second["success"]:
            return {
                "status": "Yes",
                "working_url": second["working_url"],
                "http_code": second["http_code"],
                "details": (
                    f"Original URL failed; alternate protocol worked: {alt}"
                )
            }

        return {
            "status": "No",
            "working_url": "",
            "http_code": second["http_code"] or first["http_code"],
            "details": (
                f"Original: {first['error']} | "
                f"Alternate: {second['error']}"
            )
        }

    return {
        "status": "No",
        "working_url": "",
        "http_code": first["http_code"],
        "details": first["error"]
    }


def check_all_urls():
    print("\n========== URL CHECK ==========\n")

    urls = load_urls()
    results = []

    for item in urls:

        print(f"Checking: {item['service']}")

        result = check_url(item["url"])

        result.update({
            "project": item["project"],
            "service": item["service"],
            "original_url": item["url"]
        })

        results.append(result)

        print(
            f"  Status: {result['status']} | "
            f"Working URL: {result['working_url']}"
        )

    return results


# ============================================================
# MONGODB
# ============================================================

def check_mongodb():
    print("\n========== MONGODB BACKUP ==========\n")

    config = load_key_value_file(MONGO_FILE)
    backup_date = get_backup_date()
    results = []

    for region, values in config.items():

        server = values.get("server", "").strip()
        backup_path_template = values.get(
            "backup_path",
            ""
        ).strip()

        frequency = values.get(
            "backup_frequency",
            "DAILY"
        ).strip().upper()

        if not server or not backup_path_template:
            results.append({
                "region": region,
                "server": server,
                "path": backup_path_template,
                "size": "",
                "status": "No",
                "details": "server/backup_path missing"
            })
            continue

        # US backup happens Sunday. Since we check yesterday's backup,
        # only check it when yesterday was Sunday.
        if frequency == "SUNDAY" and not was_yesterday_sunday():

            print(
                f"{region} Mongo backup skipped - "
                f"yesterday was not Sunday"
            )

            results.append({
                "region": region,
                "server": server,
                "path": "Sunday backup only",
                "size": "",
                "status": "N/A",
                "details": (
                    f"US backup checked only for Sunday backup date. "
                    f"Target backup date: {backup_date}"
                )
            })

            continue

        backup_path = backup_path_template.replace(
            "{DATE}",
            backup_date
        )

        command = (
            f'if [ -d "{backup_path}" ]; then '
            f'du -sh "{backup_path}"; '
            f'else echo "NOT_FOUND"; fi'
        )

        ssh_result = execute_ssh_command(
            server,
            command
        )

        size = ""
        status = "No"
        details = ""

        if ssh_result["success"]:

            output = ssh_result["output"]

            if output == "NOT_FOUND":
                status = "No"
                details = "Backup directory not found"

            else:
                match = re.search(
                    r"^([0-9.]+[KMGTP]?)\s+",
                    output
                )

                if match:
                    size = match.group(1)

                # Directory exists and du worked.
                # 0 size is valid.
                status = "Yes"
                details = output

        else:
            details = ssh_result["error"]

        results.append({
            "region": region,
            "server": server,
            "path": backup_path,
            "size": size,
            "status": status,
            "details": details
        })

        print(
            f"{region}: {status} | {size} | {backup_path}"
        )

    return results


# ============================================================
# RABBITMQ
# ============================================================

def check_rabbitmq():
    print("\n========== RABBITMQ ==========\n")

    config = load_key_value_file(RABBITMQ_FILE)
    results = []

    for name, values in config.items():

        server = values.get("server", "").strip()
        path = values.get("path", "").strip()
        use_sudo = values.get(
            "use_sudo",
            "false"
        ).strip().lower() == "true"

        if not server or not path:
            results.append({
                "name": name,
                "server": server,
                "path": path,
                "size": "",
                "status": "No",
                "details": "server/path missing"
            })
            continue

        # We use sudo because controller does not have direct access.
        command = f'du -sh "{path}"'

        ssh_result = execute_ssh_command(
            server,
            command,
            use_sudo=use_sudo
        )

        size = ""
        status = "No"
        details = ""

        if ssh_result["success"]:

            output = ssh_result["output"]

            match = re.search(
                r"^([0-9.]+[KMGTP]?)\s+",
                output
            )

            if match:
                size = match.group(1)
                status = "Yes"
                details = output
            else:
                details = output or "du returned no output"

        else:
            details = ssh_result["error"]

        results.append({
            "name": name,
            "server": server,
            "path": path,
            "size": size,
            "status": status,
            "details": details
        })

        print(
            f"{server}: {status} | {size}"
        )

    return results


# ============================================================
# CASSANDRA
# ============================================================

def cassandra_is_all_un(output):
    node_lines = []

    for line in output.splitlines():
        line = line.strip()

        # Cassandra status:
        # UN = Up/Normal
        # DN = Down/Normal
        # UJ = Up/Joining
        # UL = Up/Leaving
        # UM = Up/Moving
        # etc.
        if re.match(r"^[UD][NLJMD]\s+", line):
            node_lines.append(line)

    if not node_lines:
        return False, "No Cassandra node status lines found"

    all_un = all(
        line.startswith("UN ")
        for line in node_lines
    )

    if all_un:
        return True, f"All {len(node_lines)} Cassandra nodes are UN"

    non_un = [
        line.split()[0]
        for line in node_lines
        if not line.startswith("UN ")
    ]

    return False, (
        f"Non-UN node status found: "
        f"{', '.join(non_un)}"
    )


def check_cassandra():
    print("\n========== CASSANDRA ==========\n")

    config = load_key_value_file(CASSANDRA_FILE)
    results = []

    for name, values in config.items():

        server = values.get("server", "").strip()
        command = values.get(
            "command",
            "nodetool status"
        ).strip()

        ssh_result = execute_ssh_command(
            server,
            command
        )

        status = "No"
        details = ""

        if ssh_result["success"]:

            output = ssh_result["output"]

            status_bool, details = cassandra_is_all_un(
                output
            )

            status = "Yes" if status_bool else "No"

        else:
            details = ssh_result["error"]

        results.append({
            "name": name,
            "server": server,
            "command": command,
            "status": status,
            "details": details
        })

        print(
            f"{name}: {status} | {details}"
        )

    return results


# ============================================================
# EXCEL STYLES
# ============================================================

DARK_BLUE = "1F4E78"
MEDIUM_BLUE = "5B9BD5"
LIGHT_BLUE = "D9EAF7"
GREEN = "C6EFCE"
GREEN_FONT = "006100"
RED = "FFC7CE"
RED_FONT = "9C0006"
YELLOW = "FFEB9C"
YELLOW_FONT = "9C6500"
WHITE = "FFFFFF"
LIGHT_GRAY = "F2F2F2"

THIN = Side(
    style="thin",
    color="B7B7B7"
)

BORDER = Border(
    left=THIN,
    right=THIN,
    top=THIN,
    bottom=THIN
)


def apply_status_style(cell):
    value = str(cell.value or "").upper()

    if value == "YES":
        cell.fill = PatternFill(
            "solid",
            fgColor=GREEN
        )
        cell.font = Font(
            bold=True,
            color=GREEN_FONT
        )

    elif value == "NO":
        cell.fill = PatternFill(
            "solid",
            fgColor=RED
        )
        cell.font = Font(
            bold=True,
            color=RED_FONT
        )

    elif value == "N/A":
        cell.fill = PatternFill(
            "solid",
            fgColor=YELLOW
        )
        cell.font = Font(
            bold=True,
            color=YELLOW_FONT
        )


def style_section_header(ws, row, title):
    ws.merge_cells(
        start_row=row,
        start_column=1,
        end_row=row,
        end_column=7
    )

    cell = ws.cell(
        row=row,
        column=1,
        value=title
    )

    cell.fill = PatternFill(
        "solid",
        fgColor=DARK_BLUE
    )

    cell.font = Font(
        bold=True,
        color=WHITE,
        size=12
    )

    cell.alignment = Alignment(
        horizontal="left",
        vertical="center"
    )

    ws.row_dimensions[row].height = 24


def style_table_header(ws, row, headers):
    for col, header in enumerate(headers, start=1):

        cell = ws.cell(
            row=row,
            column=col,
            value=header
        )

        cell.fill = PatternFill(
            "solid",
            fgColor=MEDIUM_BLUE
        )

        cell.font = Font(
            bold=True,
            color=WHITE
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        cell.border = BORDER

    ws.row_dimensions[row].height = 30


def style_data_rows(ws, start_row, end_row, status_col):
    for row in range(start_row, end_row + 1):

        for col in range(1, 8):
            cell = ws.cell(
                row=row,
                column=col
            )

            cell.border = BORDER
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=True
            )

        apply_status_style(
            ws.cell(
                row=row,
                column=status_col
            )
        )


# ============================================================
# EXCEL GENERATION
# ============================================================

def generate_excel(
    url_results,
    mongo_results,
    rabbitmq_results,
    cassandra_results
):

    print("\n========== GENERATING EXCEL ==========\n")

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Monitoring Report"

    # --------------------------------------------------------
    # Page / sheet setup
    # --------------------------------------------------------

    ws.sheet_view.showGridLines = False

    widths = {
        "A": 24,
        "B": 42,
        "C": 24,
        "D": 70,
        "E": 70,
        "F": 14,
        "G": 16,
        "H": 16,
        "I": 60,
        "J": 24
    }

    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    # --------------------------------------------------------
    # Main title
    # --------------------------------------------------------

    ws.merge_cells("A1:J1")

    title = ws["A1"]
    title.value = "DEV INFRASTRUCTURE MONITORING REPORT"
    title.fill = PatternFill(
        "solid",
        fgColor=DARK_BLUE
    )
    title.font = Font(
        bold=True,
        color=WHITE,
        size=18
    )
    title.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:J2")

    generated = ws["A2"]
    generated.value = (
        "Generated On: "
        + datetime.now().strftime("%d-%b-%Y %H:%M:%S")
    )
    generated.font = Font(
        italic=True,
        size=10
    )
    generated.alignment = Alignment(
        horizontal="center"
    )

    current_row = 4

    # ========================================================
    # URL MONITORING
    # ========================================================

    style_section_header(
        ws,
        current_row,
        "URL MONITORING"
    )

    current_row += 1

    url_headers = [
        "Project",
        "Service",
        "Check Type",
        "URL / Path",
        "Working URL",
        "Status",
        "HTTP Code"
    ]

    style_table_header(
        ws,
        current_row,
        url_headers
    )

    current_row += 1
    url_start = current_row

    for result in url_results:

        values = [
            result["project"],
            result["service"],
            "URL",
            result["original_url"],
            result["working_url"],
            result["status"],
            result["http_code"]
        ]

        for col, value in enumerate(values, start=1):
            ws.cell(
                row=current_row,
                column=col,
                value=value
            )

        current_row += 1

    url_end = current_row - 1

    style_data_rows(
        ws,
        url_start,
        url_end,
        status_col=6
    )

    # --------------------------------------------------------
    # Merge project cells for consecutive URLs
    # --------------------------------------------------------

    if url_results:

        group_start = url_start
        previous_project = url_results[0]["project"]

        for index in range(
            1,
            len(url_results)
        ):

            current_project = url_results[index]["project"]
            actual_row = url_start + index

            if current_project != previous_project:

                if actual_row - 1 > group_start:
                    ws.merge_cells(
                        start_row=group_start,
                        start_column=1,
                        end_row=actual_row - 1,
                        end_column=1
                    )

                group_start = actual_row
                previous_project = current_project

        # Merge final project
        if url_end >= group_start:
            if url_end > group_start:
                ws.merge_cells(
                    start_row=group_start,
                    start_column=1,
                    end_row=url_end,
                    end_column=1
                )

        # Style merged project cells
        project_ranges = []

        group_start = url_start
        previous_project = url_results[0]["project"]

        for index in range(
            1,
            len(url_results)
        ):

            current_project = url_results[index]["project"]
            actual_row = url_start + index

            if current_project != previous_project:

                project_ranges.append(
                    (
                        group_start,
                        actual_row - 1
                    )
                )

                group_start = actual_row
                previous_project = current_project

        project_ranges.append(
            (
                group_start,
                url_end
            )
        )

        for start, end in project_ranges:

            cell = ws.cell(
                row=start,
                column=1
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

            cell.font = Font(
                bold=True,
                color=DARK_BLUE
            )

    current_row += 2

    # ========================================================
    # MONGODB
    # ========================================================

    style_section_header(
        ws,
        current_row,
        "MONGODB BACKUP"
    )

    current_row += 1

    mongo_headers = [
        "Region",
        "Server",
        "Check Type",
        "Backup Path",
        "Backup Date",
        "Status",
        "Size"
    ]

    style_table_header(
        ws,
        current_row,
        mongo_headers
    )

    current_row += 1
    mongo_start = current_row

    backup_date = get_backup_date()

    for result in mongo_results:

        values = [
            result["region"],
            result["server"],
            "MongoDB Backup",
            result["path"],
            backup_date,
            result["status"],
            result["size"]
        ]

        for col, value in enumerate(values, start=1):
            ws.cell(
                row=current_row,
                column=col,
                value=value
            )

        current_row += 1

    mongo_end = current_row - 1

    if mongo_results:
        style_data_rows(
            ws,
            mongo_start,
            mongo_end,
            status_col=6
        )

    current_row += 2

    # ========================================================
    # RABBITMQ
    # ========================================================

    style_section_header(
        ws,
        current_row,
        "RABBITMQ MESSAGE STORE"
    )

    current_row += 1

    rabbit_headers = [
        "Name",
        "Server",
        "Check Type",
        "Path",
        "Sudo",
        "Status",
        "Size"
    ]

    style_table_header(
        ws,
        current_row,
        rabbit_headers
    )

    current_row += 1
    rabbit_start = current_row

    rabbit_config = load_key_value_file(RABBITMQ_FILE)

    for result in rabbitmq_results:

        use_sudo = rabbit_config.get(
            result["name"],
            {}
        ).get(
            "use_sudo",
            "false"
        ).upper()

        values = [
            result["name"],
            result["server"],
            "RabbitMQ msg_stores",
            result["path"],
            use_sudo,
            result["status"],
            result["size"]
        ]

        for col, value in enumerate(values, start=1):
            ws.cell(
                row=current_row,
                column=col,
                value=value
            )

        current_row += 1

    rabbit_end = current_row - 1

    if rabbitmq_results:
        style_data_rows(
            ws,
            rabbit_start,
            rabbit_end,
            status_col=6
        )

    current_row += 2

    # ========================================================
    # CASSANDRA
    # ========================================================

    style_section_header(
        ws,
        current_row,
        "CASSANDRA"
    )

    current_row += 1

    cassandra_headers = [
        "Cluster",
        "Server",
        "Check Type",
        "Command",
        "Node Rule",
        "Status",
        "Details"
    ]

    style_table_header(
        ws,
        current_row,
        cassandra_headers
    )

    current_row += 1
    cassandra_start = current_row

    for result in cassandra_results:

        values = [
            result["name"],
            result["server"],
            "Cassandra Health",
            result["command"],
            "All nodes must be UN",
            result["status"],
            result["details"]
        ]

        for col, value in enumerate(values, start=1):
            ws.cell(
                row=current_row,
                column=col,
                value=value
            )

        current_row += 1

    cassandra_end = current_row - 1

    if cassandra_results:
        style_data_rows(
            ws,
            cassandra_start,
            cassandra_end,
            status_col=6
        )

    # --------------------------------------------------------
    # General formatting
    # --------------------------------------------------------

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A5:G{url_end}" if url_results else "A1:G1"

    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.orientation = "landscape"

    # Print area
    ws.print_area = f"A1:J{current_row - 1}"

    wb.save(REPORT_FILE)

    print(
        f"Excel report created: {REPORT_FILE}"
    )


# ============================================================
# MAIN
# ============================================================

def main():

    print()
    print("=" * 44)
    print(" DEV INFRASTRUCTURE MONITORING")
    print("=" * 44)

    url_results = check_all_urls()
    mongo_results = check_mongodb()
    rabbitmq_results = check_rabbitmq()
    cassandra_results = check_cassandra()

    generate_excel(
        url_results,
        mongo_results,
        rabbitmq_results,
        cassandra_results
    )

    print()
    print("=" * 44)
    print(" Monitoring completed successfully.")
    print("=" * 44)


if __name__ == "__main__":
    main()
